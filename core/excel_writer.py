# -*- coding: utf-8 -*-
"""Gera o arquivo .xlsx final a partir das linhas já montadas (core.builder.build_rows)."""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADERS = ["PLC", "Categoria", "Estação", "Módulo (rede)", "Endereço IP",
           "Tag (alias)", "Endereço físico", "Descrição", "Rotina(s) rastreada(s)",
           "Conflito", "Tags conflitantes"]
COL_KEYS = ["PLC", "Categoria", "Estacao", "Modulo", "IP",
            "Tag", "EnderecoFisico", "Descricao", "Rotinas",
            "Conflito", "TagsConflitantes"]
WIDTHS = [14, 14, 12, 22, 15, 26, 22, 40, 38, 10, 32]
WRAP_COLS = (8, 9, 11)   # Descrição, Rotina(s), Tags conflitantes
STATION_COL = 3

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
CELL_FONT = Font(name=FONT_NAME, size=10)
STATION_FONT = Font(name=FONT_NAME, size=10, bold=True)
CONFLICT_FONT = Font(name=FONT_NAME, size=10, bold=True, color="9C0006")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALT_FILL = PatternFill("solid", fgColor="F2F6FA")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
INFRA_FILL = PatternFill("solid", fgColor="EDEDED")
CONFLICT_FILL = PatternFill("solid", fgColor="FFC7CE")

_INVALID_SHEET_CHARS = re.compile(r'[:\\/?*\[\]]')


def safe_sheet_name(name, used_names):
    clean = _INVALID_SHEET_CHARS.sub("_", name).strip() or "PLC"
    clean = clean[:31]
    base = clean
    i = 2
    while clean.lower() in used_names:
        suffix = f"_{i}"
        clean = (base[: 31 - len(suffix)] + suffix)
        i += 1
    used_names.add(clean.lower())
    return clean


def _style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def _write_table(ws, data_rows, group_key="Estacao"):
    ws.append(HEADERS)
    _style_header(ws, len(HEADERS))
    prev_group = None
    alt = False
    r = 2
    for row in data_rows:
        gk = row.get(group_key, "") + "|" + row["PLC"]
        if gk != prev_group:
            alt = not alt
            prev_group = gk
        vals = [row[k] for k in COL_KEYS]
        ws.append(vals)
        is_conflict = row.get("Conflito") == "Sim"
        is_warn = (row["Modulo"] == "(nao localizado)") or ("multiplas estacoes" in row["Metodo"]) \
                  or (row["Metodo"] == "Nome do modulo (sem rotina)")
        is_infra = row["Categoria"] == "Infraestrutura/Outros"
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c)
            if is_conflict:
                cell.font = CONFLICT_FONT
            elif c == STATION_COL:
                cell.font = STATION_FONT
            else:
                cell.font = CELL_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(c in WRAP_COLS))
            if is_conflict:
                cell.fill = CONFLICT_FILL
            elif is_warn:
                cell.fill = WARN_FILL
            elif is_infra:
                cell.fill = INFRA_FILL if not alt else ALT_FILL
            elif alt:
                cell.fill = ALT_FILL
        r += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{r - 1}"
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


NOTES_TEXT = [
    "Notas sobre a extração da Lista de I/O",
    "",
    "Metodologia:",
    "  1. São lidos TODOS os módulos de rede (I/O remoto) configurados na árvore de I/O de cada arquivo L5X selecionado.",
    "  2. Para cada tag do tipo Alias (que aponta para o endereço físico de um módulo), é rastreado em quais rotinas",
    "     de programa (RLL) essa tag é referenciada. Alias criados no programa mas NUNCA referenciados em nenhuma",
    "     rotina são desconsiderados — não entram na lista, pois não representam um sinal realmente em uso.",
    "  3. A estação é determinada pelo nome da(s) rotina(s) encontrada(s), seguindo o padrão observado no projeto",
    "     ('L022_CIC_OP005T01', 'L301_READ_OP047T01' etc. -> estação = último trecho após '_', iniciado em 'OP').",
    "  4. Quando a tag não é referenciada em nenhuma rotina com código de estação, é usado como alternativa (fallback)",
    "     o próprio nome do módulo de rede. Essas linhas aparecem destacadas em amarelo.",
    "  5. Quando uma mesma tag aparece em rotinas de MAIS de uma estação, todas aparecem na coluna 'Estação'",
    "     separadas por ' / ', também destacado em amarelo.",
    "  6. Módulos/tags que não pertencem a nenhuma estação de produção (racks de segurança, switches Ethernet, HMIs,",
    "     quadros de comando, links de comunicação entre controladores, fontes, etc.) são agrupados como",
    "     'Infraestrutura/Outros', junto com módulos que não têm nenhuma tag alias associada.",
    "",
    "  7. 'Endereço físico' SEMPRE é um endereço de módulo (o que o time de elétrica precisa) — nunca o nome de",
    "     uma tag. Em Logix, endereço de módulo usa DOIS-PONTOS ('_005T02BM01:I1.Data[0].1'); um caminho com PONTO",
    "     ('_005T02BM01.DI01.1') é uma tag de programa, mesmo quando ela tem exatamente o mesmo nome do módulo.",
    "     Essas tags são espelhos, preenchidas por uma instrução COP/CPS/MOV dentro das rotinas de atualização de",
    "     entrada/saída (nome contendo 'ATUALIZA' ou 'AGGIORNA'). O endereço de módulo real é calculado a partir",
    "     dessa cópia, usando a POSIÇÃO do membro dentro do tipo (UDT) — e não o número no nome do membro, já que",
    "     a numeração pula (ex.: P_FESTO_VI_4832 tem DI01..DI05 e depois DI11..DI15). O bit é preservado.",
    "        exemplo: CPS(_005T02BM01:I1.Data[0], _005T02BM01.DI01, 4)",
    "                 _005T02BM01.DI01.1 -> _005T02BM01:I1.Data[0].1",
    "                 _005T02BM01.DI02.1 -> _005T02BM01:I1.Data[1].1",
    "     Quando não há cópia rastreável que cubra aquele ponto, 'Módulo' e 'Endereço físico' ficam como",
    "     '(nao localizado)', destacado em amarelo — nunca é mostrado o nome de uma tag interna nessa coluna.",
    "",
    "  8. Checagem de integridade: duas tags Alias diferentes nunca deveriam apontar para o mesmo endereço físico",
    "     (mesmo bit do mesmo módulo) — isso indica um erro de configuração no PLC (dois sinais lógicos diferentes",
    "     amarrados na mesma fiação física). Quando isso é encontrado, as linhas envolvidas são marcadas na coluna",
    "     'Conflito' e destacadas em vermelho, com a(s) outra(s) tag(s) envolvida(s) listada(s) em",
    "     'Tags conflitantes'. Essas linhas também aparecem juntas na aba 'Conflitos de Endereço'.",
    "",
    "Colunas:",
    "  • Categoria: 'Estacao' ou 'Infraestrutura/Outros'.",
    "  • Endereço físico: sempre o sinal do módulo de rede (nunca uma tag interna), resolvido automaticamente",
    "    quando necessário — ver item 7 da metodologia.",
    "  • Rotina(s) rastreada(s): lista das rotinas RLL em que a tag foi encontrada — auditoria do rastreamento.",
    "  • Conflito: 'Sim' quando o Endereço físico dessa linha também é usado por outra tag.",
    "  • Tags conflitantes: nome da(s) outra(s) tag(s) que usam o mesmo Endereço físico desta linha.",
    "",
    "Cores:",
    "  • Vermelho: conflito de endereço físico (duas tags diferentes no mesmo ponto de I/O) — prioridade sobre",
    "    qualquer outro destaque, deve ser revisado no projeto.",
    "  • Amarelo: estação ambígua (múltiplas rotinas) ou inferida só pelo nome do módulo (sem confirmação por rotina).",
    "  • Cinza: linhas de infraestrutura/outros.",
    "",
    "Limitações conhecidas:",
    "  • O rastreamento de rotina é feito por presença do nome da tag no texto (RLL) de cada rotina — cobre o caso",
    "    de instruções diretas (XIC, OTE, MOV, COP, etc.) referenciando a tag pelo nome. Referências indiretas",
    "    (endereçamento indexado, arrays acessados por variável) não são capturadas e caem no fallback por nome",
    "    do módulo, ou na categoria Infraestrutura/Outros.",
    "  • Tags cujo endereço físico ('AliasFor') não bateu com nenhum módulo conhecido aparecem com",
    "    Módulo = '(nao localizado)', destacadas em amarelo.",
]


def write_workbook(all_rows, out_path, include_infra=True):
    """
    all_rows: lista de listas de rows (uma por PLC, na ordem em que foram processados),
              cada row é um dict como produzido por core.builder.build_rows.
    """
    wb = Workbook()
    used_names = set()

    flat_rows = [row for plc_rows in all_rows for row in plc_rows]

    ws_all = wb.active
    ws_all.title = safe_sheet_name("Lista Completa", used_names)
    _write_table(ws_all, flat_rows)

    plcs_seen = []
    for row in flat_rows:
        if row["PLC"] not in plcs_seen:
            plcs_seen.append(row["PLC"])

    for plc in plcs_seen:
        ws = wb.create_sheet(safe_sheet_name(plc, used_names))
        _write_table(ws, [r for r in flat_rows if r["PLC"] == plc and r["Categoria"] == "Estacao"])

    if include_infra:
        ws_infra = wb.create_sheet(safe_sheet_name("Infraestrutura e Outros", used_names))
        _write_table(ws_infra, [r for r in flat_rows if r["Categoria"] != "Estacao"], group_key="Modulo")

    conflitos = sorted(
        (r for r in flat_rows if r.get("Conflito") == "Sim"),
        key=lambda r: (r["PLC"], r["EnderecoFisico"], r["Tag"]),
    )
    if conflitos:
        ws_conf = wb.create_sheet(safe_sheet_name("Conflitos de Endereço", used_names))
        _write_table(ws_conf, conflitos, group_key="EnderecoFisico")

    ws_notes = wb.create_sheet(safe_sheet_name("Notas", used_names))
    ws_notes.column_dimensions["A"].width = 105
    for i, line in enumerate(NOTES_TEXT, start=1):
        cell = ws_notes.cell(row=i, column=1, value=line)
        cell.font = Font(name=FONT_NAME, size=11, bold=(i == 1))
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(out_path)

    n_estacao = sum(1 for r in flat_rows if r["Categoria"] == "Estacao")
    n_infra = sum(1 for r in flat_rows if r["Categoria"] != "Estacao")
    return {
        "total_linhas": len(flat_rows),
        "linhas_estacao": n_estacao,
        "linhas_infra": n_infra,
        "linhas_conflito": len(conflitos),
        "plcs": len(plcs_seen),
    }
